"""Ambil teks dari isi file (PDF, Word, Excel, TXT/CSV, opsional OCR foto) untuk pencarian pintar.

Hasilnya disimpan di kolom documents.content_text (teks lengkap) dan content_tokens
(daftar kata/kode unik — ringkas, disalin ke aplikasi supaya pencarian isi file
tetap jalan saat offline).
"""
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

MAX_TEXT_CHARS = 2_000_000      # teks lengkap maksimum yang disimpan per dokumen
MAX_TOKENS_CHARS = 60_000       # daftar token maksimum per dokumen (dikirim ke aplikasi)

STOPWORDS = {
    "dan", "yang", "di", "ke", "dari", "untuk", "dengan", "pada", "ini", "itu", "atau", "dalam", "akan", "oleh",
    "adalah", "sebagai", "tidak", "ada", "juga", "telah", "tersebut", "bahwa", "serta", "kepada", "agar", "bagi",
    "the", "of", "and", "to", "in", "for", "on", "is", "by", "with", "as", "at", "an", "be", "this", "that",
}

_CODE_RE = re.compile(r"\d+(?:[./\-]\d+)+")         # 5.2.2.01  600/123  01-02-2026
_WORD_RE = re.compile(r"[0-9a-zA-ZÀ-ɏ]+")


def build_tokens(*texts: str) -> str:
    """Kata & kode unik (huruf kecil), dipisah spasi. Kode bertitik juga disimpan versi rapatnya
    (5.2.2.01 -> 52201) supaya cocok diketik dengan atau tanpa titik."""
    seen: dict[str, None] = {}
    total = 0
    for text in texts:
        if not text:
            continue
        low = text.lower()
        for m in _CODE_RE.finditer(low):
            code = m.group(0)
            for t in (code, re.sub(r"[./\-]", "", code)):
                if t not in seen:
                    seen[t] = None
                    total += len(t) + 1
        for m in _WORD_RE.finditer(low):
            t = m.group(0)
            if len(t) < 2 or t in STOPWORDS or t in seen:
                continue
            seen[t] = None
            total += len(t) + 1
            if total > MAX_TOKENS_CHARS:
                break
        if total > MAX_TOKENS_CHARS:
            break
    return " ".join(seen.keys())


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    size = 0
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        parts.append(t)
        size += len(t)
        if size > MAX_TEXT_CHARS:
            break
    return "\n".join(parts)


def _docx(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        texts = []
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml") and (
                "document" in name or "header" in name or "footer" in name or "footnotes" in name
            ):
                root = ElementTree.fromstring(z.read(name))
                for para in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                    texts.append("".join(t.text or "" for t in para.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")))
        return "\n".join(texts)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    size = 0
    for ws in wb.worksheets:
        out.append(f"[{ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None and str(v).strip()]
            if vals:
                line = " | ".join(vals)
                out.append(line)
                size += len(line)
                if size > MAX_TEXT_CHARS:
                    break
    wb.close()
    return "\n".join(out)


def _text(path: Path) -> str:
    raw = path.read_bytes()[:MAX_TEXT_CHARS]
    for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="ignore")


def _ocr(path: Path) -> str:
    """OCR foto/scan — hanya jalan kalau pytesseract + Tesseract-OCR terpasang (opsional)."""
    import pytesseract  # noqa: F401  (ImportError -> dianggap tidak didukung)
    from PIL import Image

    return pytesseract.image_to_string(Image.open(path), lang="ind+eng")


def extract_text(path: Path, ext: str) -> tuple[str, str]:
    """-> (teks, status). status: ok | empty | unsupported | error:<pesan>"""
    ext = ext.lower()
    try:
        if ext == "pdf":
            text = _pdf(path)
        elif ext == "docx":
            text = _docx(path)
        elif ext == "xlsx":
            text = _xlsx(path)
        elif ext in ("txt", "csv"):
            text = _text(path)
        elif ext in ("jpg", "jpeg", "png", "bmp", "webp"):
            try:
                text = _ocr(path)
            except ImportError:
                return "", "unsupported"
        else:
            return "", "unsupported"  # doc/xls lama, dwg, zip, rar
    except Exception as e:  # file rusak / terkunci password
        return "", f"error:{type(e).__name__}"
    text = (text or "")[:MAX_TEXT_CHARS]
    return text, ("ok" if text.strip() else "empty")


def snippet(text: str, terms: list[str], width: int = 90) -> str:
    """Potongan teks di sekitar kata yang cocok (untuk ditampilkan di hasil pencarian)."""
    low = text.lower()
    for t in terms:
        i = low.find(t.lower())
        if i >= 0:
            a = max(0, i - width)
            b = min(len(text), i + len(t) + width)
            s = re.sub(r"\s+", " ", text[a:b]).strip()
            return ("…" if a > 0 else "") + s + ("…" if b < len(text) else "")
    return ""
