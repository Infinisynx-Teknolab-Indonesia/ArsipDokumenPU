"""Data contoh untuk testing: 9 proyek, 2 akun staf, dan 50 dokumen ASLI (PDF, Word,
Excel, TXT, foto PNG) dari tahun 2000 sampai 2026.

Sebagian dokumen SALING BERHUBUNGAN supaya fitur pencarian pintar & dokumen terkait
bisa langsung dicoba, sebagian lagi BERDIRI SENDIRI (tidak ada relasinya):

  A. Seri "Dokumen 1": Formula A (2000) -> Formula B (2002) -> Formula C (2026)
  B. Kontrak 2100/KTR/PU/2000 Jalan Lingkar Sekupang -> SPMK, Addendum, BA PHO/FHO,
     kontrak Tahap II 2002 (semua menyebut nomor 2100)
  C. Kode rekening 544 (pengeluaran/belanja) di isi file 2010, 2015, 2019 (5.4.4), 2025, 2026
     -> cari "pengeluaran" ikut menemukan dokumen berisi 544 (lewat Kamus)
  D. Seri "AHSP Bidang Jalan" 2015 -> Revisi 1 2019 -> 2025
  E. Jembatan Sei Ladi lama (BM-0458, 2010) -> jembatan baru 2026 (menyebut BM-0458)
  F. Gedung Kantor Kecamatan Sekupang 2015 (kontrak 3307) -> rehabilitasi 2025
  G. Drainase Ahmad Yani 2019 (saluran DR-1120) -> normalisasi 2026
  H. Proyek berjalan Simpang Kara 2026: kontrak, laporan mingguan, MC, foto progres
  I. 10 dokumen TANPA relasi (surat, SOP, pedoman, notulen, inventaris, dll.)

Dijalankan otomatis saat server pertama kali start (SEED_DEMO=1 di .env) atau manual:
    D:\\python3.12\\python.exe seed_demo.py
"""
from __future__ import annotations

import hashlib
import io
import re
import struct
import zipfile
import zlib
from datetime import date, datetime, timedelta

from sqlalchemy import func, select

from . import config
from .db import Document, Project, SessionLocal, User
from .security import make_password_fields

HEADER = ["PEMERINTAH KOTA BATAM", "DINAS PEKERJAAN UMUM DAN PENATAAN RUANG", "Jl. Engku Putri No. 1, Batam Centre", ""]

# --------------------------------------------------------------------------- proyek & akun

PROJECTS = [
    dict(code="PU-JLN-2000-003", name="Pembangunan Jalan Lingkar Sekupang", type="Jalan", location="Kec. Sekupang (STA 0+000 – 4+500)",
         year=2000, funding_source="APBD Kabupaten/Kota", contract_value=3_250_000_000, contractor="PT Karya Lingkar Utama",
         supervisor="CV Mitra Teknik", ppk="Ir. Syamsul Bahri", contract_number="2100/KTR/PU/2000",
         start_date=date(2000, 4, 5), end_date=date(2000, 12, 15), status="FHO / Selesai", progress=100),
    dict(code="PU-JLN-2002-011", name="Pelebaran Jalan Lingkar Sekupang Tahap II", type="Jalan", location="Kec. Sekupang (STA 4+500 – 7+000)",
         year=2002, funding_source="APBD Kabupaten/Kota", contract_value=2_100_000_000, contractor="PT Karya Lingkar Utama",
         supervisor="CV Mitra Teknik", ppk="Ir. Syamsul Bahri", contract_number="2100-II/KTR/PU/2002",
         start_date=date(2002, 3, 10), end_date=date(2002, 11, 30), status="FHO / Selesai", progress=100),
    dict(code="PU-JBT-2010-004", name="Pembangunan Jembatan Sei Ladi (Lama)", type="Jembatan", location="Kel. Kibing, Batu Aji",
         year=2010, funding_source="APBN", contract_value=6_800_000_000, contractor="PT Beton Nusantara", supervisor="PT Rekayasa Mandiri",
         ppk="Drs. Hendra Gunawan", contract_number="0458/KTR/BM/2010", start_date=date(2010, 3, 1), end_date=date(2010, 12, 10),
         status="FHO / Selesai", progress=100),
    dict(code="PU-GDG-2015-002", name="Pembangunan Gedung Kantor Kecamatan Sekupang", type="Gedung / Bangunan", location="Kec. Sekupang",
         year=2015, funding_source="APBD Kabupaten/Kota", contract_value=4_400_000_000, contractor="CV Bangun Sejahtera",
         supervisor="CV Arsitek Muda", ppk="Ahmad Fauzi, S.T.", contract_number="3307/KTR/CK/2015",
         start_date=date(2015, 3, 2), end_date=date(2015, 11, 30), status="FHO / Selesai", progress=100),
    dict(code="PU-DRN-2019-006", name="Drainase Primer Jl. Ahmad Yani Tahap I", type="Drainase", location="Batam Kota",
         year=2019, funding_source="APBD Provinsi", contract_value=1_950_000_000, contractor="PT Tirta Utama", supervisor="CV Hidro Konsultan",
         ppk="Ir. Rahmat Hidayat", contract_number="1120/KTR/SDA/2019", start_date=date(2019, 3, 1), end_date=date(2019, 11, 20),
         status="FHO / Selesai", progress=100),
    dict(code="PU-JLN-2026-001", name="Peningkatan Jalan Simpang Kara – Batu Aji", type="Jalan", location="Kec. Batu Aji (STA 0+000 – 3+200)",
         year=2026, funding_source="APBD Kabupaten/Kota", contract_value=4_850_000_000, contractor="PT Karya Jalan Nusantara",
         supervisor="CV Cipta Konsultan", ppk="Ir. Rahmat Hidayat", contract_number="600/KTR/PUPR-BM/2026/014",
         start_date=date(2026, 3, 2), end_date=date(2026, 11, 28), status="Pelaksanaan", progress=62),
    dict(code="PU-JBT-2026-002", name="Pembangunan Jembatan Sei Ladi", type="Jembatan", location="Kel. Kibing, Batu Aji",
         year=2026, funding_source="DAK", contract_value=12_300_000_000, contractor="PT Beton Perkasa", supervisor="PT Rekayasa Mandiri",
         ppk="Dewi Lestari, S.T., M.T.", contract_number="600/KTR/PUPR-BM/2026/021", start_date=date(2026, 4, 15),
         end_date=date(2027, 2, 10), status="Pelaksanaan", progress=35),
    dict(code="PU-GDG-2025-007", name="Rehabilitasi Gedung Kantor Kecamatan Sekupang", type="Gedung / Bangunan", location="Kec. Sekupang",
         year=2025, funding_source="APBD Kabupaten/Kota", contract_value=1_975_000_000, contractor="CV Bangun Sejahtera",
         supervisor="CV Arsitek Muda", ppk="Ahmad Fauzi, S.T.", contract_number="600/KTR/PUPR-CK/2025/007",
         start_date=date(2025, 5, 5), end_date=date(2025, 10, 30), status="Masa Pemeliharaan", progress=100),
    dict(code="PU-DRN-2026-004", name="Normalisasi Drainase Primer Jl. Ahmad Yani", type="Drainase", location="Batam Kota",
         year=2026, funding_source="APBD Provinsi", contract_value=2_400_000_000, contractor="PT Tirta Utama", supervisor="CV Hidro Konsultan",
         ppk="Ir. Rahmat Hidayat", contract_number="", status="Lelang / Pengadaan", progress=0),
]

USERS = [
    dict(username="budi", full_name="Budi Santoso", role="staff", nip="198705122010011003", jabatan="Pengawas Lapangan", password="budi12345"),
    dict(username="sari", full_name="Sari Wulandari", role="staff", nip="199102032015032001", jabatan="Staf Administrasi Teknik", password="sari12345"),
]

# --------------------------------------------------------------------------- 50 dokumen
# (judul, nomor, tanggal, kategori, kode proyek, tag, seri, jenis file, isi [baris])

D = date

DOCS: list[tuple] = [
    # A. Seri "Dokumen 1" — Formula A / B / C
    ("Dokumen 1 Formula A", "2100/FRM/PU/2000", D(2000, 3, 15), "RAB", "PU-JLN-2000-003", "formula, eskalasi, harga", "Dokumen 1", "pdf",
     ["FORMULA PENYESUAIAN HARGA (ESKALASI) - FORMULA A",
      "P(n) = 0,15 + 0,40 Bt/Bo + 0,30 Mt/Mo + 0,15 Et/Eo",
      "Bt = indeks harga bahan aspal, Mt = indeks material agregat, Et = indeks peralatan",
      "Berlaku untuk kontrak nomor 2100/KTR/PU/2000 Jalan Lingkar Sekupang.",
      "Indeks dasar (Bo) mengacu BPS bulan Januari 2000."]),
    ("Dokumen 1 Formula B", "2100/FRM/PU/2002", D(2002, 6, 10), "RAB", "PU-JLN-2002-011", "formula, eskalasi, harga", "Dokumen 1", "pdf",
     ["FORMULA PENYESUAIAN HARGA (ESKALASI) - FORMULA B (perubahan dari Formula A tahun 2000)",
      "P(n) = 0,15 + 0,35 Bt/Bo + 0,30 Mt/Mo + 0,10 Et/Eo + 0,10 Lt/Lo",
      "Ditambahkan komponen upah tenaga kerja (Lt).",
      "Berlaku untuk kontrak lanjutan 2100-II/KTR/PU/2002."]),
    ("Dokumen 1 Formula C", "600/FRM/PUPR/2026/003", D(2026, 2, 20), "RAB", "PU-JLN-2026-001", "formula, eskalasi, harga", "Dokumen 1", "docx",
     ["FORMULA PENYESUAIAN HARGA (ESKALASI) - FORMULA C (pengganti Formula A 2000 dan Formula B 2002)",
      "P(n) = 0,15 + 0,30 Bt/Bo + 0,25 Mt/Mo + 0,15 Et/Eo + 0,15 Lt/Lo",
      "Indeks mengacu Permen PUPR tentang penyesuaian harga terbaru.",
      "Dipakai untuk paket Peningkatan Jalan Simpang Kara 2026."]),

    # B. Kontrak 2100 Jalan Lingkar Sekupang
    ("Kontrak Pembangunan Jalan Lingkar Sekupang", "2100/KTR/PU/2000", D(2000, 4, 1), "KONTRAK", "PU-JLN-2000-003", "kontrak, aspal", "", "pdf",
     ["SURAT PERJANJIAN (KONTRAK) NOMOR 2100/KTR/PU/2000",
      "Pekerjaan: Pembangunan Jalan Lingkar Sekupang STA 0+000 - 4+500",
      "Nilai kontrak: Rp 3.250.000.000,- (termasuk PPN)",
      "Penyedia: PT Karya Lingkar Utama",
      "Jangka waktu pelaksanaan 240 hari kalender, masa pemeliharaan 180 hari."]),
    ("SPMK Jalan Lingkar Sekupang", "2101/SPMK/PU/2000", D(2000, 4, 5), "SPMK", "PU-JLN-2000-003", "spmk", "", "pdf",
     ["SURAT PERINTAH MULAI KERJA",
      "Berdasarkan kontrak nomor 2100/KTR/PU/2000 tanggal 1 April 2000,",
      "penyedia diperintahkan memulai pekerjaan Pembangunan Jalan Lingkar Sekupang",
      "terhitung sejak tanggal 5 April 2000."]),
    ("Addendum I Kontrak Jalan Lingkar Sekupang", "2100-A1/KTR/PU/2000", D(2000, 9, 12), "KONTRAK", "PU-JLN-2000-003", "addendum, kontrak", "", "docx",
     ["ADDENDUM I ATAS KONTRAK NOMOR 2100/KTR/PU/2000",
      "Perubahan volume pekerjaan lapis aspal AC-WC dari 4.500 m2 menjadi 4.950 m2.",
      "Nilai kontrak berubah menjadi Rp 3.410.000.000,-",
      "Ketentuan lain dalam kontrak 2100 tetap berlaku."]),
    ("BA PHO Jalan Lingkar Sekupang", "2150/BA-PHO/PU/2000", D(2000, 12, 20), "BA", "PU-JLN-2000-003", "pho, serah terima", "", "pdf",
     ["BERITA ACARA SERAH TERIMA PERTAMA PEKERJAAN (PHO)",
      "Kontrak nomor 2100/KTR/PU/2000 beserta Addendum I.",
      "Pekerjaan dinyatakan selesai 100% dan diterima dengan catatan perbaikan minor marka jalan."]),
    ("BA FHO Jalan Lingkar Sekupang", "2190/BA-FHO/PU/2001", D(2001, 6, 22), "BA", "PU-JLN-2000-003", "fho, serah terima", "", "pdf",
     ["BERITA ACARA SERAH TERIMA AKHIR PEKERJAAN (FHO)",
      "Masa pemeliharaan kontrak 2100/KTR/PU/2000 telah berakhir.",
      "Jaminan pemeliharaan dapat dikembalikan kepada penyedia."]),
    ("Kontrak Pelebaran Jalan Lingkar Sekupang Tahap II", "2100-II/KTR/PU/2002", D(2002, 3, 1), "KONTRAK", "PU-JLN-2002-011", "kontrak, pelebaran", "", "pdf",
     ["SURAT PERJANJIAN NOMOR 2100-II/KTR/PU/2002",
      "Pekerjaan lanjutan dari kontrak 2100/KTR/PU/2000 (Jalan Lingkar Sekupang).",
      "Pelebaran STA 4+500 - 7+000 menjadi 2 x 7 meter.",
      "Nilai kontrak Rp 2.100.000.000,-"]),

    # C. Kode rekening 544 (pengeluaran)
    ("SPJ Belanja Pemeliharaan Jalan Triwulan I", "900/SPJ/BM/2010/01", D(2010, 4, 10), "KEUANGAN", "", "spj, belanja", "", "xlsx",
     [["Kode Rekening", "Uraian", "Jumlah (Rp)"],
      ["544", "Belanja pemeliharaan jalan dan jembatan", 185_000_000],
      ["544.01", "Tambal sulam aspal", 64_500_000],
      ["544.02", "Pembersihan saluran tepi jalan", 21_300_000]]),
    ("Rekap Pengeluaran Kegiatan Gedung Kecamatan", "900/RKP/CK/2015/12", D(2015, 12, 15), "KEUANGAN", "PU-GDG-2015-002", "rekap, realisasi", "", "xlsx",
     [["Kode", "Uraian", "Pagu", "Realisasi"],
      ["544", "Belanja modal gedung kantor", 4_400_000_000, 4_312_750_000],
      ["522", "Belanja jasa konsultansi pengawasan", 132_000_000, 131_500_000]]),
    ("Laporan Realisasi Anggaran Bidang Bina Marga", "900/LRA/BM/2019", D(2019, 12, 30), "KEUANGAN", "", "lra, anggaran", "", "pdf",
     ["LAPORAN REALISASI ANGGARAN TAHUN 2019 - BIDANG BINA MARGA",
      "5.4.4  Belanja Modal Jalan, Irigasi dan Jaringan   Rp 18.450.000.000   realisasi 96,2%",
      "5.2.2  Belanja Barang dan Jasa                     Rp  2.115.000.000   realisasi 91,4%",
      "Sisa anggaran dikembalikan ke kas daerah."]),
    ("Bukti Pembayaran Termin 2 Jembatan Sei Ladi", "SP2D-0773/2026", D(2026, 7, 10), "KEUANGAN", "PU-JBT-2026-002", "termin, sp2d", "", "pdf",
     ["SURAT PERINTAH PENCAIRAN DANA (SP2D) NOMOR 0773/2026",
      "Pembayaran termin 2 (progres fisik 30%) Pembangunan Jembatan Sei Ladi",
      "Dibebankan pada kode rekening 544 Belanja Modal Jembatan.",
      "Jumlah dibayarkan Rp 2.952.000.000,-"]),
    ("Nota Dinas Pengeluaran Uang Muka", "ND-118/PUPR/2025", D(2025, 6, 2), "SURAT", "PU-GDG-2025-007", "uang muka, nota dinas", "", "txt",
     ["NOTA DINAS",
      "Perihal: Permohonan pengeluaran uang muka 20% Rehabilitasi Gedung Kantor Kecamatan Sekupang",
      "Mohon diproses pada kode rekening 544 sesuai DPA tahun anggaran 2025.",
      "Jaminan uang muka dari bank telah diverifikasi."]),

    # D. Seri AHSP
    ("Analisa Harga Satuan Pekerjaan (AHSP) Bidang Jalan", "050/AHSP/BM/2015", D(2015, 1, 20), "RAB", "", "ahsp, harga satuan", "AHSP Bidang Jalan", "xlsx",
     [["Kode", "Uraian Pekerjaan", "Satuan", "Harga Satuan (Rp)"],
      ["7.1", "Lapis Aus Aspal (AC-WC)", "ton", 1_245_000], ["7.2", "Lapis Antara (AC-BC)", "ton", 1_198_000],
      ["5.1", "Lapis Pondasi Agregat Kelas A", "m3", 498_000]]),
    ("Analisa Harga Satuan Pekerjaan (AHSP) Bidang Jalan Revisi 1", "050/AHSP/BM/2019", D(2019, 2, 11), "RAB", "", "ahsp, harga satuan", "AHSP Bidang Jalan", "xlsx",
     [["Kode", "Uraian Pekerjaan", "Satuan", "Harga Satuan (Rp)"],
      ["7.1", "Lapis Aus Aspal (AC-WC)", "ton", 1_410_000], ["7.2", "Lapis Antara (AC-BC)", "ton", 1_356_000],
      ["5.1", "Lapis Pondasi Agregat Kelas A", "m3", 562_000]]),
    ("Analisa Harga Satuan Pekerjaan (AHSP) Bidang Jalan", "050/AHSP/BM/2025", D(2025, 1, 15), "RAB", "PU-JLN-2026-001", "ahsp, harga satuan", "AHSP Bidang Jalan", "xlsx",
     [["Kode", "Uraian Pekerjaan", "Satuan", "Harga Satuan (Rp)"],
      ["7.1", "Lapis Aus Aspal (AC-WC)", "ton", 1_685_000], ["7.2", "Lapis Antara (AC-BC)", "ton", 1_602_000],
      ["5.1", "Lapis Pondasi Agregat Kelas A", "m3", 671_000], ["9.3", "Beton fc' 30 MPa", "m3", 1_950_000]]),

    # E. Jembatan Sei Ladi lama (BM-0458) -> baru 2026
    ("Uji Beton Jembatan Sei Ladi K-300", "LAB-0458/2010", D(2010, 8, 14), "UJI_MUTU", "PU-JBT-2010-004", "uji beton, k-300", "", "pdf",
     ["HASIL UJI KUAT TEKAN BETON - JEMBATAN SEI LADI (BM-0458)",
      "Mutu rencana K-300. Benda uji kubus 15x15x15 cm umur 28 hari.",
      "Rata-rata kuat tekan 318 kg/cm2 - MEMENUHI SYARAT."]),
    ("BA PHO Jembatan Sei Ladi", "0459/BA-PHO/BM/2010", D(2010, 12, 18), "BA", "PU-JBT-2010-004", "pho, jembatan", "", "docx",
     ["BERITA ACARA SERAH TERIMA PERTAMA - JEMBATAN SEI LADI",
      "Kontrak 0458/KTR/BM/2010. Nomor aset jembatan BM-0458.",
      "Bentang 30 m, lebar 9 m, struktur girder beton bertulang."]),
    ("Gambar As-Built Jembatan Sei Ladi", "ASB-0458/2011", D(2011, 3, 1), "ASBUILT", "PU-JBT-2010-004", "as-built, gambar", "", "pdf",
     ["GAMBAR TERLAKSANA (AS-BUILT DRAWING) JEMBATAN SEI LADI - BM-0458",
      "Lembar 1: Denah dan potongan memanjang, bentang 30 m",
      "Lembar 2: Detail abutment A1 dan A2, pondasi tiang pancang 40x40 cm",
      "Elevasi lantai jembatan +4,250 dari muka air rencana."]),
    ("Kontrak Pembangunan Jembatan Sei Ladi", "600/KTR/PUPR-BM/2026/021", D(2026, 4, 10), "KONTRAK", "PU-JBT-2026-002", "kontrak, jembatan", "", "pdf",
     ["SURAT PERJANJIAN NOMOR 600/KTR/PUPR-BM/2026/021",
      "Pembangunan Jembatan Sei Ladi pengganti jembatan lama nomor aset BM-0458 (dibangun 2010).",
      "Bentang baru 45 m, girder PCI, lebar 11 m.",
      "Nilai kontrak Rp 12.300.000.000,- sumber dana DAK."]),
    ("Shop Drawing Girder Jembatan Sei Ladi", "SD-021/2026", D(2026, 5, 20), "GAMBAR", "PU-JBT-2026-002", "shop drawing, girder", "", "pdf",
     ["SHOP DRAWING GIRDER PCI H-210 JEMBATAN SEI LADI",
      "Posisi abutment baru bergeser 7,5 m dari abutment lama BM-0458.",
      "Lihat gambar as-built 2011 untuk posisi tiang pancang lama."]),
    ("Laporan Uji Tiang Pancang Jembatan Sei Ladi", "LAB-021/2026", D(2026, 6, 15), "UJI_MUTU", "PU-JBT-2026-002", "pda test, tiang pancang", "", "txt",
     ["LAPORAN PILE DRIVING ANALYZER (PDA) TEST",
      "Lokasi: Abutment A1 Jembatan Sei Ladi (pengganti BM-0458)",
      "Daya dukung ultimit tiang uji 1: 182 ton (rencana 150 ton) - OK",
      "Daya dukung ultimit tiang uji 2: 176 ton (rencana 150 ton) - OK"]),

    # F. Gedung kantor kecamatan
    ("Kontrak Pembangunan Gedung Kantor Kecamatan Sekupang", "3307/KTR/CK/2015", D(2015, 3, 1), "KONTRAK", "PU-GDG-2015-002", "kontrak, gedung", "", "pdf",
     ["SURAT PERJANJIAN NOMOR 3307/KTR/CK/2015",
      "Pembangunan Gedung Kantor Kecamatan Sekupang 2 lantai, luas 1.200 m2.",
      "Nilai kontrak Rp 4.400.000.000,-"]),
    ("Gambar As-Built Gedung Kantor Kecamatan Sekupang", "ASB-3307/2016", D(2016, 1, 10), "ASBUILT", "PU-GDG-2015-002", "as-built, gedung", "", "pdf",
     ["GAMBAR TERLAKSANA GEDUNG KANTOR KECAMATAN SEKUPANG (kontrak 3307)",
      "Denah lantai 1 dan 2, tampak, potongan, detail atap baja ringan."]),
    ("Kontrak Rehabilitasi Gedung Kantor Kecamatan Sekupang", "600/KTR/PUPR-CK/2025/007", D(2025, 5, 2), "KONTRAK", "PU-GDG-2025-007", "kontrak, rehabilitasi", "", "pdf",
     ["SURAT PERJANJIAN NOMOR 600/KTR/PUPR-CK/2025/007",
      "Rehabilitasi gedung hasil kontrak 3307/KTR/CK/2015: penggantian atap, plafon dan instalasi listrik.",
      "Nilai kontrak Rp 1.975.000.000,-"]),
    ("Foto Progres 100% Rehabilitasi Gedung Kecamatan", "", D(2025, 10, 28), "FOTO", "PU-GDG-2025-007", "foto, progres 100%", "", "png", []),

    # G. Drainase Ahmad Yani
    ("Gambar Rencana Saluran Ahmad Yani", "GR-1120/2019", D(2019, 3, 5), "GAMBAR", "PU-DRN-2019-006", "gambar rencana, saluran", "", "pdf",
     ["GAMBAR RENCANA SALURAN PRIMER DR-1120 JL. AHMAD YANI",
      "Box culvert 2,0 x 1,5 m sepanjang 850 m.",
      "Kemiringan memanjang 0,2%."]),
    ("Laporan Akhir Drainase Primer Jl. Ahmad Yani Tahap I", "LA-1120/2019", D(2019, 12, 10), "LAP_BULANAN", "PU-DRN-2019-006", "laporan akhir, drainase", "", "docx",
     ["LAPORAN AKHIR PEKERJAAN DRAINASE PRIMER JL. AHMAD YANI TAHAP I",
      "Saluran DR-1120 selesai 850 m, sisa 620 m direncanakan tahap berikutnya.",
      "Genangan di simpang Ahmad Yani berkurang dari 60 cm menjadi 10 cm."]),
    ("HPS Normalisasi Drainase Ahmad Yani", "HPS-1120/2026", D(2026, 2, 1), "RAB", "PU-DRN-2026-004", "hps, drainase", "", "xlsx",
     [["No", "Uraian", "Volume", "Satuan", "Harga (Rp)"],
      [1, "Normalisasi saluran DR-1120 (lanjutan tahap I 2019)", 620, "m", 1_850_000_000],
      [2, "Pengerukan sedimen", 1450, "m3", 290_000_000]]),
    ("Izin Lingkungan Normalisasi Drainase", "660/IL/DLH/2026/045", D(2026, 3, 11), "IZIN", "PU-DRN-2026-004", "izin, lingkungan", "", "pdf",
     ["PERSETUJUAN LINGKUNGAN (UKL-UPL)",
      "Kegiatan: Normalisasi Drainase Primer Jl. Ahmad Yani",
      "Pemrakarsa wajib mengelola sedimen hasil pengerukan ke lokasi yang ditentukan."]),

    # H. Proyek berjalan Simpang Kara 2026
    ("Kontrak Peningkatan Jalan Simpang Kara", "600/KTR/PUPR-BM/2026/014", D(2026, 3, 2), "KONTRAK", "PU-JLN-2026-001", "kontrak, aspal", "", "pdf",
     ["SURAT PERJANJIAN NOMOR 600/KTR/PUPR-BM/2026/014",
      "Peningkatan Jalan Simpang Kara - Batu Aji STA 0+000 - 3+200",
      "Nilai kontrak Rp 4.850.000.000,- penyesuaian harga memakai Formula C."]),
    ("MC-0 Peningkatan Jalan Simpang Kara", "MC0-014/2026", D(2026, 3, 5), "MC", "PU-JLN-2026-001", "mc0, mutual check", "", "xlsx",
     [["Item", "Uraian", "Volume Kontrak", "Volume MC-0", "Satuan"],
      ["7.1", "Lapis Aus AC-WC", 2350, 2410, "ton"], ["5.1", "Agregat Kelas A", 1800, 1765, "m3"]]),
    ("Foto Progres 0% Simpang Kara", "", D(2026, 3, 6), "FOTO", "PU-JLN-2026-001", "foto, progres 0%", "", "png", []),
    ("Laporan Mingguan ke-1 Peningkatan Jalan Simpang Kara", "LM-01/014/2026", D(2026, 3, 9), "LAP_MINGGUAN", "PU-JLN-2026-001", "laporan mingguan, minggu 1", "", "pdf",
     ["LAPORAN MINGGUAN KE-1", "Mobilisasi alat berat dan pengukuran ulang (uitzet).", "Progres rencana 2,1% - realisasi 2,4%."]),
    ("Laporan Mingguan ke-2 Peningkatan Jalan Simpang Kara", "LM-02/014/2026", D(2026, 3, 16), "LAP_MINGGUAN", "PU-JLN-2026-001", "laporan mingguan, minggu 2", "", "docx",
     ["LAPORAN MINGGUAN KE-2", "Galian badan jalan STA 0+000 - 0+400.", "Progres rencana 4,8% - realisasi 4,5%."]),
    ("Laporan Mingguan ke-3 Peningkatan Jalan Simpang Kara", "LM-03/014/2026", D(2026, 3, 23), "LAP_MINGGUAN", "PU-JLN-2026-001", "laporan mingguan, minggu 3", "", "pdf",
     ["LAPORAN MINGGUAN KE-3", "Penghamparan agregat kelas A STA 0+000 - 0+350.", "Hujan 3 hari, progres rencana 7,9% - realisasi 7,1%."]),
    ("Laporan Mingguan ke-4 Peningkatan Jalan Simpang Kara", "LM-04/014/2026", D(2026, 3, 30), "LAP_MINGGUAN", "PU-JLN-2026-001", "laporan mingguan, minggu 4", "", "txt",
     ["LAPORAN MINGGUAN KE-4", "Pemadatan agregat dan uji sand cone.", "Progres rencana 11,2% - realisasi 11,0%."]),
    ("Foto Progres 50% Simpang Kara", "", D(2026, 7, 2), "FOTO", "PU-JLN-2026-001", "foto, progres 50%", "", "png", []),
    ("MC-50 Peningkatan Jalan Simpang Kara", "MC50-014/2026", D(2026, 7, 1), "MC", "PU-JLN-2026-001", "mc50, opname", "", "xlsx",
     [["Item", "Uraian", "Volume s.d. lalu", "Volume periode ini", "Satuan"],
      ["7.1", "Lapis Aus AC-WC", 0, 1180, "ton"], ["5.1", "Agregat Kelas A", 1765, 0, "m3"]]),

    # I. Dokumen TANPA relasi
    ("Surat Undangan Rapat Koordinasi Bulanan", "005/UND/SEKR/2005", D(2005, 2, 14), "SURAT", "", "undangan, rapat", "", "txt",
     ["UNDANGAN", "Rapat koordinasi bulanan kepala bidang.", "Hari Senin, 21 Februari 2005, pukul 09.00 di ruang rapat lantai 2."]),
    ("Pedoman Pengarsipan Dokumen Teknis", "PED-01/SEKR/2008", D(2008, 11, 20), "LAINNYA", "", "pedoman, arsip", "", "docx",
     ["PEDOMAN PENGARSIPAN DOKUMEN TEKNIS", "Setiap dokumen diberi kode klasifikasi, disimpan minimal 10 tahun.",
      "Dokumen kontrak disimpan permanen."]),
    ("SOP Keselamatan Kerja di Lapangan", "SOP-K3/07/2012", D(2012, 7, 1), "K3", "", "sop, k3, apd", "", "pdf",
     ["STANDAR OPERASIONAL PROSEDUR K3 LAPANGAN", "Wajib helm, rompi reflektif, sepatu safety.",
      "Rambu pengalihan lalu lintas dipasang 50 m sebelum area kerja."]),
    ("Daftar Inventaris Alat Berat", "INV-AB/2017", D(2017, 4, 12), "LAINNYA", "", "inventaris, alat berat", "", "xlsx",
     [["No", "Nama Alat", "Merk", "Tahun", "Kondisi"], [1, "Excavator", "Komatsu PC200", 2011, "Baik"],
      [2, "Vibro Roller", "Sakai SV512", 2014, "Baik"], [3, "Dump Truck", "Hino 500", 2016, "Rusak ringan"]]),
    ("Sertifikat Kalibrasi Alat Ukur Theodolit", "KAL-TH/0918", D(2018, 9, 5), "UJI_MUTU", "", "kalibrasi, theodolit", "", "pdf",
     ["SERTIFIKAT KALIBRASI", "Alat: Theodolit digital Topcon DT-209, nomor seri 7Q8812.", "Hasil: dalam batas toleransi 5 detik."]),
    ("Peta Rawan Banjir Kota", "", D(2020, 1, 15), "GAMBAR", "", "peta, banjir", "", "png", []),
    ("Notulen Rapat Evaluasi Kinerja Bidang", "NOT-08/SEKR/2021", D(2021, 8, 30), "SURAT", "", "notulen, evaluasi", "", "docx",
     ["NOTULEN RAPAT EVALUASI KINERJA", "Capaian indikator kinerja bidang semester I sebesar 87%.",
      "Tindak lanjut: percepatan pengadaan paket fisik."]),
    ("Laporan Tahunan Dinas", "LT/PUPR/2022", D(2022, 12, 31), "LAP_BULANAN", "", "laporan tahunan", "", "pdf",
     ["LAPORAN TAHUNAN DINAS", "Ringkasan capaian program dan kegiatan selama satu tahun anggaran.",
      "Panjang jalan kondisi mantap meningkat menjadi 78,4%."]),
    ("Surat Edaran Jam Kerja Bulan Ramadhan", "SE-03/SEKR/2023", D(2023, 3, 20), "SURAT", "", "surat edaran, jam kerja", "", "txt",
     ["SURAT EDARAN", "Selama bulan Ramadhan jam kerja pukul 08.00 - 15.00 (Senin - Kamis).", "Jumat pukul 08.00 - 15.30."]),
    ("Rencana Strategis Bidang Bina Marga", "RENSTRA/BM/2024", D(2024, 2, 1), "LAINNYA", "", "renstra, perencanaan", "", "pdf",
     ["RENCANA STRATEGIS BIDANG BINA MARGA", "Sasaran: kemantapan jalan kota 90% di akhir periode.",
      "Strategi: preservasi rutin, rekonstruksi ruas kritis, digitalisasi arsip teknis."]),
]

assert len(DOCS) == 50, len(DOCS)

# --------------------------------------------------------------------------- pembuat file


def _ascii(s: str) -> str:
    return (s.replace("–", "-").replace("—", "-").replace("’", "'").replace("“", '"').replace("”", '"')
            .encode("latin-1", "replace").decode("latin-1"))


def make_pdf(lines: list[str]) -> bytes:
    """PDF sederhana (Helvetica) — teksnya bisa dibaca pypdf untuk pencarian isi file."""
    def esc(t: str) -> str:
        return _ascii(t).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    body = ["BT", "/F1 11 Tf", "15 TL", "50 790 Td"]
    for i, line in enumerate(lines):
        if i == 3:
            body.append("/F1 12 Tf")
        body.append(f"({esc(line)}) Tj T*")
    body.append("ET")
    stream = "\n".join(body).encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(lines: list[str]) -> bytes:
    """Dokumen Word (.docx) minimal yang bisa dibuka Microsoft Word."""
    from xml.sax.saxutils import escape
    paras = "".join(
        f'<w:p><w:r>{"<w:rPr><w:b/></w:rPr>" if i in (0, 1, 4) else ""}<w:t xml:space="preserve">{escape(t)}</w:t></w:r></w:p>'
        for i, t in enumerate(lines)
    )
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                f'{paras}<w:sectPr/></w:body></w:document>')
    content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                     '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '</Relationships>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
    return buf.getvalue()


def make_xlsx(title: str, header: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for line in HEADER[:2]:
        ws.append([line])
    ws.append([title])
    ws.append([])
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for col, width in zip("ABCDE", (14, 48, 18, 18, 12)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_png(seed: int, w: int = 640, h: int = 400) -> bytes:
    """'Foto' contoh: langit, tanah, dan jalan (gambar sederhana tanpa library tambahan)."""
    raw = bytearray()
    sky = (110 + seed * 7 % 60, 160 + seed * 3 % 50, 220)
    for y in range(h):
        raw.append(0)
        for x in range(w):
            if y < h * 0.45:
                c = (sky[0] + y // 8, sky[1] + y // 10, min(255, sky[2]))
            elif abs((x - w / 2) - (y - h * 0.45) * 0.0) < (y - h * 0.45) * 1.4 + 20:
                stripe = (y // 18) % 2 == 0 and abs(x - w / 2) < 4
                c = (250, 250, 250) if stripe else (70, 72, 78)
            else:
                c = (90 + seed * 5 % 40, 140, 70)
            raw.extend(bytes(c))

    def chunk(t: bytes, d: bytes) -> bytes:
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(bytes(raw), 6)) + chunk(b"IEND", b""))


def _safe_name(title: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")[:80]


def build_file(i: int, title: str, number: str, d: date, project: Project | None, kind: str, content: list) -> tuple[str, bytes]:
    name = f"{_safe_name(title)}.{kind}"
    meta = [f"Nomor : {number or '-'}", f"Tanggal : {d.strftime('%d-%m-%Y')}"]
    if project:
        meta.append(f"Paket : {project.code} - {project.name}")
    if kind == "pdf":
        return name, make_pdf(HEADER + meta + [""] + content)
    if kind == "docx":
        return name, make_docx(HEADER + meta + [""] + content)
    if kind == "txt":
        return name, ("\r\n".join(HEADER + meta + [""] + content)).encode("utf-8")
    if kind == "xlsx":
        return name, make_xlsx(title, content[0], content[1:])
    if kind == "png":
        return name, make_png(i)
    raise ValueError(kind)


# --------------------------------------------------------------------------- seeding


def seed(with_documents: bool = True, verbose: bool = True) -> dict:
    """Isi proyek, akun staf, dan 50 dokumen contoh. Aman diulang (yang sudah ada dilewati)."""
    from .main import index_document  # import di sini supaya tidak melingkar

    added = {"projects": 0, "users": 0, "documents": 0}
    with SessionLocal() as db:
        for p in PROJECTS:
            if not db.scalar(select(Project).where(Project.code == p["code"])):
                db.add(Project(**p))
                added["projects"] += 1
        for u in USERS:
            if not db.scalar(select(User).where(User.username == u["username"])):
                data = dict(u)
                pw = data.pop("password")
                db.add(User(active=True, **data, **make_password_fields(pw)))
                added["users"] += 1
        db.commit()

        if not with_documents:
            return added

        projects = {p.code: p for p in db.scalars(select(Project)).all()}
        admin = db.scalar(select(User.username).where(User.role == "admin").order_by(User.id)) or "admin"
        uploaders = [admin, "budi", "sari"]
        new_ids = []
        for i, (title, number, d, cat, pcode, tags, series, kind, content) in enumerate(DOCS, 1):
            cid = f"demo-{i:02d}"
            if db.scalar(select(Document.id).where(Document.client_id == cid)):
                continue
            proj = projects.get(pcode) if pcode else None
            fname, data = build_file(i, title, number, d, proj, kind, content)
            created = datetime(d.year, d.month, d.day, 9, 0) + timedelta(minutes=i * 7)
            folder = config.STORAGE_DIR / f"{created:%Y}" / f"{created:%m}"
            folder.mkdir(parents=True, exist_ok=True)
            target = folder / f"{cid}_{fname}"
            target.write_bytes(data)
            doc = Document(
                client_id=cid, title=title, doc_number=number, doc_date=d, category=cat,
                project_id=proj.id if proj else None,
                description=(content[0] if content and isinstance(content[0], str) else f"Data contoh {title}")[:500],
                tags=tags, series=series, file_name=fname,
                stored_path=str(target.relative_to(config.STORAGE_DIR)), file_size=len(data),
                mime={"pdf": "application/pdf", "txt": "text/plain", "png": "image/png",
                      "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                      "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}[kind],
                sha256=hashlib.sha256(data).hexdigest(), uploaded_by=uploaders[i % 3],
                created_at=created, uploaded_at=created + timedelta(minutes=3), content_status="pending",
            )
            db.add(doc)
            db.flush()
            new_ids.append(doc.id)
            added["documents"] += 1
        db.commit()

    for doc_id in new_ids:  # baca isi file sekarang supaya pencarian isi langsung jalan
        try:
            index_document(doc_id)
        except Exception as e:  # pragma: no cover
            print(f"[demo] gagal index #{doc_id}: {e}")
    if verbose:
        print(f"[demo] data contoh ditambahkan: {added['projects']} proyek, {added['users']} akun staf, {added['documents']} dokumen")
    return added


def documents_empty() -> bool:
    with SessionLocal() as db:
        return (db.scalar(select(func.count()).select_from(Document)) or 0) == 0
